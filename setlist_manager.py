import os
import shutil
import string
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, simpledialog

from openpyxl import load_workbook


class SetlistManager:

    def __init__(self, root):
        self.root = root
        self.root.title("MIDI Setlist Manager")
        self.root.geometry("1200x600")

        self.source_folder = ""
        self.max_rounds = 9  # 8 Runden + Restliche Lieder
        self.round_trees = {}

        self.all_midi_files = set()
        self.song_counts = {}
        self.used_midi_files = set()

        # Drag & Drop
        self.dragged_item = None

        self.create_ui()

    # ---------------- UI ----------------

    def create_ui(self):
        main_frame = tk.Frame(self.root)
        main_frame.pack(fill="both", expand=True, padx=10, pady=10)

        # LINKS – Songs
        left_frame = tk.Frame(main_frame)
        left_frame.pack(side="left", fill="both", expand=True)

        self.left_list = tk.Listbox(left_frame, selectmode=tk.SINGLE)
        self.left_list.pack(fill="both", expand=True)

        tk.Button(left_frame, text="📂 Ordner wählen", command=self.load_folder).pack(pady=5)

        # MITTE – Buttons
        middle_frame = tk.Frame(main_frame)
        middle_frame.pack(side="left", padx=10)

        tk.Button(middle_frame, text="➜ Hinzufügen", width=14, command=self.add_to_round).pack(pady=5)
        tk.Button(middle_frame, text="⬅ Entfernen", width=14, command=self.remove_from_round).pack(pady=5)
        tk.Button(middle_frame, text="💾 Export", width=14, command=self.export_setlist).pack(pady=20)

        # RECHTS – Runden
        right_frame = tk.Frame(main_frame)
        right_frame.pack(side="left", fill="both", expand=True)

        self.notebook = ttk.Notebook(right_frame)
        self.notebook.pack(fill="both", expand=True)

        for i in range(1, self.max_rounds + 1):
            tab = tk.Frame(self.notebook)
            name = f"Runde {i}" if i < self.max_rounds else "Restliche Lieder"
            self.notebook.add(tab, text=name)

            lf = tk.LabelFrame(tab, text=name)
            lf.pack(fill="both", expand=True, padx=5, pady=5)

            tree = ttk.Treeview(
                lf,
                columns=("num", "song"),
                show="headings",
                selectmode="browse",
                height=20
            )
            tree.heading("num", text="#" if i < self.max_rounds else "")
            tree.column("num", width=40, anchor="center")
            tree.heading("song", text="Songname")
            tree.column("song", width=400, anchor="w")
            tree.pack(fill="both", expand=True)

            tree.bind("<ButtonPress-1>", self.on_start_drag)
            tree.bind("<B1-Motion>", self.on_drag)
            tree.bind("<ButtonRelease-1>", self.on_drop)

            self.round_trees[i] = tree

    # ---------------- Logic ----------------

    def load_folder(self):
        folder = filedialog.askdirectory()
        if not folder:
            return

        self.source_folder = folder
        self.left_list.delete(0, tk.END)

        files = sorted(f for f in os.listdir(folder) if f.lower().endswith((".mid", ".midi")))
        self.all_midi_files = set(files)
        self.song_counts = {f: 0 for f in files}  # reset counts

        for f in files:
            self.left_list.insert(tk.END, f)

        self.rebuild_rest_tree()

    def populate_restliche_lieder(self):
        all_files = set(os.listdir(self.source_folder))
        used_files = set()

        for rnd in range(1, self.max_rounds):
            tree = self.round_trees[rnd]
            for iid in tree.get_children():
                used_files.add(tree.set(iid, "song"))

        remaining_files = sorted(all_files - used_files)

        tree_rest = self.round_trees[self.max_rounds]
        for f in remaining_files:
            if f.lower().endswith((".mid", ".midi")):
                tree_rest.insert("", "end", values=("", f))

    def get_current_round(self):
        return self.notebook.index(self.notebook.select()) + 1

    def update_numbers(self, rnd):
        if rnd == self.max_rounds:
            return
        tree = self.round_trees[rnd]
        for i, iid in enumerate(tree.get_children(), start=1):
            tree.set(iid, "num", i)

    def add_to_round(self):
        sel = self.left_list.curselection()
        if not sel:
            return

        file = self.left_list.get(sel)
        rnd = self.get_current_round()
        if rnd == self.max_rounds:
            return

        if file in self.used_midi_files:
            return

        tree = self.round_trees[rnd]
        tree.insert("", "end", values=("", file))
        self.update_numbers(rnd)

        self.used_midi_files.add(file)
        self.rebuild_rest_tree()

    def remove_from_round(self):
        rnd = self.get_current_round()

        if rnd == self.max_rounds:
            return

        tree = self.round_trees[rnd]

        removed = []
        for iid in tree.selection():
            removed.append(tree.set(iid, "song"))
            tree.delete(iid)

        self.update_numbers(rnd)

        for file in removed:
            self.used_midi_files.discard(file)

        self.rebuild_rest_tree()

    # ---------------- Drag & Drop ----------------

    def on_start_drag(self, event):
        tree = event.widget
        iid = tree.identify_row(event.y)
        if iid:
            self.dragged_item = iid
            tree.selection_set(iid)

    def on_drag(self, event):
        tree = event.widget
        iid = tree.identify_row(event.y)
        if iid:
            tree.selection_set(iid)

    def on_drop(self, event):
        tree = event.widget
        if not self.dragged_item:
            return

        target = tree.identify_row(event.y)
        if target:
            tree.move(self.dragged_item, "", tree.index(target))
            self.update_numbers(self.get_current_round())

        self.dragged_item = None

    # ---------------- Export ----------------

    def get_suffix(self, idx):
        letters = list(string.ascii_lowercase)
        return letters[idx] if idx < 26 else "z" * (idx - 25 + 1)

    def export_excel_from_template(self, export_folder):
        template = os.path.join(self.source_folder, "Setlist_Vorlage.xlsx")
        if not os.path.exists(template):
            messagebox.showerror("Fehler", "Setlist_Vorlage.xlsx nicht gefunden!")
            return

        titel = simpledialog.askstring("Titel", "Bitte Titel für die Setlist eingeben:")
        if titel is None:
            return

        wb = load_workbook(template)
        ws = wb.active

        # Titel
        ws["A1"] = titel
        ws["K1"] = titel

        start_col = 2  # B
        header_row = 2  # Runde XX
        first_song_row = 3  # Songs

        for rnd in range(1, self.max_rounds):
            tree = self.round_trees[rnd]

            col = start_col + (rnd - 1) * 2
            col_letter = ws.cell(row=1, column=col).column_letter

            ws[f"{col_letter}{header_row}"] = f"Runde {rnd:02d}"

            # Alte Songs löschen
            for r in range(first_song_row, 200):
                ws[f"{col_letter}{r}"] = ""

            # Songs schreiben
            for i, iid in enumerate(tree.get_children()):
                ws[f"{col_letter}{first_song_row + i}"] = tree.set(iid, "song")

        wb.save(os.path.join(export_folder, "Setlist.xlsx"))

    def export_setlist(self):
        export_folder = filedialog.askdirectory(title="Export Ordner wählen")
        if not export_folder:
            return

        target = os.path.join(export_folder, "Setlist_Export")
        os.makedirs(target, exist_ok=True)

        for rnd in range(1, self.max_rounds + 1):
            tree = self.round_trees[rnd]
            for i, iid in enumerate(tree.get_children()):
                name = tree.set(iid, "song")
                sourcepath = os.path.join(self.source_folder, name)

                if rnd == self.max_rounds:
                    new_name = name
                else:
                    new_name = f"{rnd}{self.get_suffix(i)}.{name}"

                dest_path = os.path.join(export_folder, new_name)
                shutil.copy2(sourcepath, os.path.join(target, new_name))

        self.export_excel_from_template(target)

        messagebox.showinfo("Fertig", "Export abgeschlossen!")

    def rebuild_rest_tree(self):
        tree_rest = self.round_trees[self.max_rounds]
        for iid in tree_rest.get_children():
            tree_rest.delete(iid)

        remaining = sorted(self.all_midi_files - self.used_midi_files)
        for f in remaining:
            tree_rest.insert("", "end", values=("", f))

if __name__ == "__main__":
    root = tk.Tk()
    app = SetlistManager(root)
    root.mainloop()
